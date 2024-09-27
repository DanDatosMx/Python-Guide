#Llenado de Información de un Excel a un Word
#Para llenar información de un Excel a un documento Word, puedes usar openpyxl para Excel y python-docx para Word:

import openpyxl
from docx import Document

# Cargar datos desde Excel
wb = openpyxl.load_workbook('datos.xlsx')
sheet = wb.active

# Crear un documento Word
doc = Document()

# Llenar el documento con datos de Excel
for row in sheet.iter_rows(min_row=2, values_only=True):
    doc.add_paragraph(f'Nombre: {row[0]}, Edad: {row[1]}, Ciudad: {row[2]}')

# Guardar el documento Word
doc.save('informe.docx')
